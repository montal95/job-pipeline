"""
Writer agent tests — CV loading, gap extraction, prompt builders, JSON parsers,
python-docx rendering, node behavior, and deferred-render contract.

All tests mock the Anthropic API — no live LLM calls.
The two deferred-render regression tests (previously in test_phase4) live here
because they test Writer agent behavior, not the Submitter.
"""

from __future__ import annotations

import json
import sqlite3
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from pydantic import ValidationError

FIXTURES = Path(__file__).parent / "fixtures"
SAMPLE_CV = FIXTURES / "sample_cv.txt"
SAMPLE_RESUME_LLM = FIXTURES / "sample_resume_llm_response.json"
SAMPLE_CL_LLM = FIXTURES / "sample_cover_letter_llm_response.json"


@pytest.fixture()
def cv_text() -> str:
    return SAMPLE_CV.read_text(encoding="utf-8")


@pytest.fixture()
def writer_state(cv_text, sample_job):
    """
    empty_state() pre-loaded with the standard test CV, job shortlist, and
    empty interview answers. Used by LLM node tests to avoid repeating setup.
    """
    from pipeline.state import empty_state
    state = empty_state()
    state["cv_text"] = cv_text
    state["shortlist"] = [sample_job]
    state["current_job_id"] = sample_job.id
    state["interview_answers"] = {}
    return state


# ── CV loading ─────────────────────────────────────────────────────────────────


def test_load_cv_text_reads_file():
    from pipeline.agents.writer import _load_cv_text
    result = _load_cv_text(str(SAMPLE_CV))
    assert isinstance(result, str)
    assert len(result) > 100
    assert "Samuel Montalvo" in result


def test_load_cv_text_missing_file_raises():
    from pipeline.agents.writer import _load_cv_text
    with pytest.raises(FileNotFoundError):
        _load_cv_text("/nonexistent/path/cv.txt")


# ── Gap extraction ─────────────────────────────────────────────────────────────


def test_extract_job_gaps_finds_missing_skills(cv_text):
    from pipeline.agents.writer import _extract_job_gaps
    gaps = _extract_job_gaps(cv_text, "Experience with Kubernetes and Terraform required.")
    assert any("kubernetes" in g.lower() for g in gaps)
    assert any("terraform" in g.lower() for g in gaps)


def test_extract_job_gaps_no_gaps(cv_text):
    from pipeline.agents.writer import _extract_job_gaps
    gaps = _extract_job_gaps(cv_text, "Experience with Rails and PostgreSQL required.")
    assert gaps == []


def test_extract_job_gaps_case_insensitive(cv_text):
    from pipeline.agents.writer import _extract_job_gaps
    gaps = _extract_job_gaps(cv_text, "Strong postgresql skills required.")
    assert not any("postgresql" in g.lower() for g in gaps)


# ── Prompt builders ────────────────────────────────────────────────────────────


def test_build_resume_prompt_contains_job_title(cv_text, sample_job):
    from pipeline.agents.writer import _build_resume_prompt
    prompt = _build_resume_prompt(cv_text, sample_job, answers={})
    assert "Senior Software Engineer" in prompt


def test_build_cover_letter_prompt_contains_company(cv_text, sample_job):
    from pipeline.agents.writer import _build_cover_letter_prompt
    prompt = _build_cover_letter_prompt(cv_text, sample_job, answers={})
    assert "Acme Health" in prompt


# ── JSON parsers ───────────────────────────────────────────────────────────────


def test_parse_resume_json_happy_path():
    from pipeline.agents.writer import _parse_resume_json
    from pipeline.state import ResumeContent
    result = _parse_resume_json(SAMPLE_RESUME_LLM.read_text(encoding="utf-8"))
    assert isinstance(result, ResumeContent)
    assert result.name == "Samuel Montalvo"
    assert len(result.sections) > 0
    assert len(result.skills) > 0


def test_parse_resume_json_missing_field_raises():
    from pipeline.agents.writer import _parse_resume_json
    with pytest.raises(ValidationError):
        _parse_resume_json(json.dumps({"name": "Test"}))


def test_parse_cover_letter_json_happy_path():
    from pipeline.agents.writer import _parse_cover_letter_json
    from pipeline.state import CoverLetterContent
    result = _parse_cover_letter_json(SAMPLE_CL_LLM.read_text(encoding="utf-8"))
    assert isinstance(result, CoverLetterContent)
    assert len(result.opening) > 0
    assert len(result.body_paragraphs) > 0
    assert len(result.closing) > 0


def test_parse_cover_letter_json_missing_field_raises():
    from pipeline.agents.writer import _parse_cover_letter_json
    with pytest.raises(ValidationError):
        _parse_cover_letter_json(json.dumps({"opening": "Hello"}))


# ── docx rendering ─────────────────────────────────────────────────────────────


def test_render_resume_docx_creates_file(resume_content, tmp_path):
    from pipeline.agents.writer import _render_resume_docx
    out = str(tmp_path / "test_resume.docx")
    assert Path(_render_resume_docx(resume_content, out)).exists()


def test_render_resume_docx_contains_name(resume_content, tmp_path):
    from docx import Document
    from pipeline.agents.writer import _render_resume_docx
    out = str(tmp_path / "test_resume.docx")
    _render_resume_docx(resume_content, out)
    full_text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "Samuel Montalvo" in full_text


def test_render_cover_letter_docx_creates_file(cover_letter_content, tmp_path):
    from pipeline.agents.writer import _render_cover_letter_docx
    out = str(tmp_path / "test_cover_letter.docx")
    assert Path(_render_cover_letter_docx(cover_letter_content, out)).exists()


def test_render_cover_letter_docx_contains_opening(cover_letter_content, tmp_path):
    from docx import Document
    from pipeline.agents.writer import _render_cover_letter_docx
    out = str(tmp_path / "test_cover_letter.docx")
    _render_cover_letter_docx(cover_letter_content, out)
    full_text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "Dear Acme Health" in full_text


# ── Node behavior ──────────────────────────────────────────────────────────────


def test_write_resume_makes_one_llm_call_returns_content(writer_state, mock_llm_message, tmp_path):
    """write_resume makes exactly one Anthropic API call and returns resume_content (no path)."""
    from pipeline.agents.writer import write_resume

    with patch("pipeline.agents.writer.anthropic.Anthropic") as mock_cls:
        mock_client = MagicMock()
        mock_cls.return_value = mock_client
        mock_client.messages.create.return_value = mock_llm_message
        with patch("pipeline.agents.writer.settings") as mock_settings:
            mock_settings.output_dir = str(tmp_path)
            mock_settings.max_revision_rounds = 3
            mock_settings.llm_provider = "anthropic"
            mock_settings.anthropic_api_key = "test-key"
            result = write_resume(writer_state)

    mock_client.messages.create.assert_called_once()
    assert result.get("resume_content") is not None
    assert result.get("resume_path") is None  # deferred-render contract


def test_apply_feedback_increments_revision_round(writer_state, mock_llm_message, tmp_path):
    from pipeline.agents.writer import apply_feedback

    writer_state["revision_round"] = 0
    writer_state["human_feedback"] = "Please make the summary more concise."

    with patch("pipeline.agents.writer.anthropic.Anthropic") as mock_cls:
        mock_client = MagicMock()
        mock_cls.return_value = mock_client
        mock_client.messages.create.return_value = mock_llm_message
        with patch("pipeline.agents.writer.settings") as mock_settings:
            mock_settings.output_dir = str(tmp_path)
            mock_settings.max_revision_rounds = 3
            mock_settings.llm_provider = "anthropic"
            mock_settings.anthropic_api_key = "test-key"
            result = apply_feedback(writer_state)

    assert result["revision_round"] == 1


def test_apply_feedback_revision_prompt_includes_formatting_rules(writer_state, mock_llm_message, tmp_path):
    """Revision prompt must include structural formatting rules so the LLM doesn't revert to markdown."""
    from pipeline.agents.writer import apply_feedback

    writer_state["revision_round"] = 0
    writer_state["human_feedback"] = "Make the summary shorter."

    with patch("pipeline.agents.writer.anthropic.Anthropic") as mock_cls:
        mock_client = MagicMock()
        mock_cls.return_value = mock_client
        mock_client.messages.create.return_value = mock_llm_message
        with patch("pipeline.agents.writer.settings") as mock_settings:
            mock_settings.output_dir = str(tmp_path)
            mock_settings.max_revision_rounds = 3
            mock_settings.llm_provider = "anthropic"
            mock_settings.anthropic_api_key = "test-key"
            apply_feedback(writer_state)

    call_args = mock_client.messages.create.call_args
    prompt = call_args[1]["messages"][0]["content"]
    assert "No markdown" in prompt
    assert "Company · Location | Role Title" in prompt
    assert "Stack:" in prompt


def test_apply_cl_feedback_increments_cl_revision_round(writer_state, mock_cl_llm_message, tmp_path):
    from pipeline.agents.writer import apply_cl_feedback
    from pipeline.state import CoverLetterContent

    writer_state["cl_revision_round"] = 0
    writer_state["human_feedback"] = "Make the opening stronger."
    writer_state["cover_letter_content"] = CoverLetterContent(
        opening="Dear Hiring Team,",
        body_paragraphs=["I have experience."],
        closing="Thank you.",
    )

    with patch("pipeline.agents.writer.anthropic.Anthropic") as mock_cls:
        mock_client = MagicMock()
        mock_cls.return_value = mock_client
        mock_client.messages.create.return_value = mock_cl_llm_message
        with patch("pipeline.agents.writer.settings") as mock_settings:
            mock_settings.llm_provider = "anthropic"
            mock_settings.anthropic_api_key = "test-key"
            mock_settings.max_revision_rounds = 3
            result = apply_cl_feedback(writer_state)

    assert result["cl_revision_round"] == 1
    assert result["human_feedback"] is None
    assert result.get("cover_letter_content") is not None


def test_should_revise_routes_correctly():
    from pipeline.agents.writer import should_revise_resume
    from pipeline.state import empty_state

    with patch("pipeline.agents.writer.settings") as mock_settings:
        mock_settings.max_revision_rounds = 3

        # feedback + rounds remaining → revise resume
        s = empty_state()
        s["human_feedback"] = "Make it shorter."
        s["revision_round"] = 1
        assert should_revise_resume(s) == "apply_feedback"

        # approved, default flow → go to CL review
        s2 = empty_state()
        s2["human_feedback"] = None
        s2["revision_round"] = 0
        assert should_revise_resume(s2) == "cl_review_interrupt"

        # approved, --resume-only → skip CL, persist
        s3 = empty_state()
        s3["human_feedback"] = None
        s3["revision_round"] = 0
        s3["write_resume_only"] = True
        assert should_revise_resume(s3) == "persist_documents"

        # max rounds hit → warn
        s4 = empty_state()
        s4["human_feedback"] = "Still not right."
        s4["revision_round"] = 3
        assert should_revise_resume(s4) == "warn_and_exit"


def test_should_revise_cl_routes_correctly():
    from pipeline.agents.writer import should_revise_cl
    from pipeline.state import empty_state

    with patch("pipeline.agents.writer.settings") as mock_settings:
        mock_settings.max_revision_rounds = 3

        s = empty_state()
        s["human_feedback"] = "More concise please."
        s["cl_revision_round"] = 1
        assert should_revise_cl(s) == "apply_cl_feedback"

        s2 = empty_state()
        s2["human_feedback"] = None
        s2["cl_revision_round"] = 0
        assert should_revise_cl(s2) == "persist_documents"

        s3 = empty_state()
        s3["human_feedback"] = "Still bad."
        s3["cl_revision_round"] = 3
        assert should_revise_cl(s3) == "warn_and_exit"


def test_should_write_resume_routes_correctly():
    from pipeline.agents.writer import should_write_resume
    from pipeline.state import empty_state

    s = empty_state()
    assert should_write_resume(s) == "write_resume"

    s2 = empty_state()
    s2["write_cover_letter_only"] = True
    assert should_write_resume(s2) == "write_cover_letter"


def test_should_write_cl_routes_correctly():
    from pipeline.agents.writer import should_write_cl
    from pipeline.state import empty_state

    s = empty_state()
    assert should_write_cl(s) == "write_cover_letter"

    s2 = empty_state()
    s2["write_resume_only"] = True
    assert should_write_cl(s2) == "persist_documents"


def test_write_resume_passes_api_key_to_anthropic(writer_state, mock_llm_message, tmp_path):
    """write_resume must instantiate Anthropic with api_key from settings, not rely on env var."""
    from pipeline.agents.writer import write_resume

    with patch("pipeline.agents.writer.anthropic.Anthropic") as mock_cls:
        mock_client = MagicMock()
        mock_cls.return_value = mock_client
        mock_client.messages.create.return_value = mock_llm_message
        with patch("pipeline.agents.writer.settings") as mock_settings:
            mock_settings.output_dir = str(tmp_path)
            mock_settings.max_revision_rounds = 3
            mock_settings.anthropic_api_key = "sk-ant-test-key"
            mock_settings.llm_provider = "anthropic"
            write_resume(writer_state)

    mock_cls.assert_called_once_with(api_key="sk-ant-test-key")


def test_pre_write_interview_no_gaps_returns_early(sample_job):
    """When gap analysis finds no missing skills, return immediately without calling interrupt()."""
    from pipeline.agents.writer import pre_write_interview
    from pipeline.state import empty_state
    from unittest.mock import patch as _patch

    # Give the job a description that overlaps fully with the CV skills
    sample_job.description = "Ruby on Rails experience required."
    state = empty_state()
    state["cv_text"] = "Experienced with Ruby on Rails."
    state["shortlist"] = [sample_job]
    state["current_job_id"] = sample_job.id

    with _patch("pipeline.agents.writer.interrupt") as mock_interrupt:
        result = pre_write_interview(state)

    mock_interrupt.assert_not_called()
    assert result == {"interview_answers": {}}


def test_pre_write_interview_with_gaps_calls_interrupt(sample_job, monkeypatch):
    """When gaps exist, interrupt() is called with questions for each gap."""
    from pipeline.agents.writer import pre_write_interview
    from pipeline.state import empty_state
    from unittest.mock import patch as _patch

    sample_job.description = "Kubernetes and Terraform experience required."
    state = empty_state()
    state["cv_text"] = "Experienced with Ruby on Rails and PostgreSQL."
    state["shortlist"] = [sample_job]
    state["current_job_id"] = sample_job.id

    captured = {}
    def fake_interrupt(payload):
        captured.update(payload)
        return {}

    with _patch("pipeline.agents.writer.interrupt", side_effect=fake_interrupt):
        pre_write_interview(state)

    assert "questions" in captured
    assert len(captured["questions"]) > 0


# ── Deferred-render contract (regressions) ────────────────────────────────────


def test_persist_documents_saves_content_json_not_paths(
    sample_job, resume_content, cover_letter_content, seeded_db
):
    """persist_documents writes content JSON to DB; file path columns stay NULL."""
    from pipeline.agents.writer import persist_documents
    from pipeline.state import empty_state

    state = empty_state()
    state["current_job_id"] = sample_job.id
    state["resume_content"] = resume_content
    state["cover_letter_content"] = cover_letter_content

    with patch("pipeline.agents.writer.settings") as mock_settings:
        mock_settings.app_db_path = seeded_db
        persist_documents(state)

    conn = sqlite3.connect(seeded_db)
    row = conn.execute(
        "SELECT resume_path, resume_content_json FROM jobs WHERE id = ?",
        (sample_job.id,),
    ).fetchone()
    conn.close()

    assert row[0] is None, "resume_path should stay NULL — deferred render"
    assert row[1] is not None, "resume_content_json should be populated"
    assert json.loads(row[1])["name"] == resume_content.name
